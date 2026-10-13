# Copyright (c) 2026, Riddhi and contributors
# For license information, please see license.txt

"""Seasonal Supplier Quantities — what to order from each supplier for the season.

Replaces the client's "Seasonal Supplier Quantities" workbook, whose tabs were
the same item list cut different ways. Here each tab is a View:

  All Locations      every submitted order line
  Pick-Up Locations  orders with a pick-up location that is not Home Delivery
                     (narrow to one location with the Pick-Up Location filter —
                     that is the per-location tab of the workbook)
  Home Delivery      orders whose pick-up location is the Home Delivery one
  Totals             quantity and weight per vendor per location

The "Download Workbook" button rebuilds the whole multi-tab workbook in one go.

Quantity is the line's stock qty and Total Weight is ERPNext's own line
total_weight (stock qty x weight per unit), so catch-weight typed on the order
line is what gets summed — not the item master's nominal weight. Weight is shown
as Total Weight / Quantity for that reason.

Only items with Maintain Stock ticked are counted — untick it on fee/credit
items (Processing Fee, FREE …) to keep them off the supplier sheet.

Vendor is the supplier on the item's Purchasing tab (Item Supplier table) —
one per item by agreement; if a second row slips in, the first row wins.
Items without one sort last and total under "No Vendor".
"""

import re
from collections import OrderedDict, defaultdict
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import flt

ALL = "All Locations"
PICKUP = "Pick-Up Locations"
HOME = "Home Delivery"
TOTALS = "Totals"


def execute(filters=None):
	filters = frappe._dict(filters or {})
	view = filters.view or ALL

	if view == TOTALS:
		return get_totals_columns(), get_totals_data(filters)

	return get_columns(), get_item_data(filters, view, filters.pick_up_location)


def get_columns():
	return [
		{"label": _("Description"), "fieldname": "item_name", "fieldtype": "Data", "width": 380},
		{"label": _("Product Code"), "fieldname": "item_code", "fieldtype": "Link", "options": "Item", "width": 120},
		{"label": _("Quantity"), "fieldname": "qty", "fieldtype": "Float", "precision": 2, "width": 90},
		{"label": _("Weight"), "fieldname": "weight", "fieldtype": "Float", "precision": 2, "width": 90},
		{"label": _("Total Weight"), "fieldname": "total_weight", "fieldtype": "Float", "precision": 2, "width": 110},
		{"label": _("Vendor"), "fieldname": "vendor", "fieldtype": "Link", "options": "Supplier", "width": 160},
		{"label": _("Sales Order"), "fieldname": "sales_order", "fieldtype": "Data", "width": 260},
		{"label": _("Pick-Up Location"), "fieldname": "pick_up_location", "fieldtype": "Data", "width": 220},
	]


def get_totals_columns():
	return [
		{"label": _("Vendor"), "fieldname": "vendor", "fieldtype": "Data", "width": 160},
		{"label": _("Location"), "fieldname": "location", "fieldtype": "Data", "width": 260},
		{"label": _("Total Qty"), "fieldname": "qty", "fieldtype": "Float", "precision": 2, "width": 110},
		{"label": _("Total Weight"), "fieldname": "total_weight", "fieldtype": "Float", "precision": 2, "width": 120},
	]


NO_VENDOR = "No Vendor"


def get_vendor_map(item_codes):
	"""item_code -> supplier from the item's Suppliers table (first row)."""
	if not item_codes:
		return {}
	vendors = {}
	for item_code, supplier in frappe.db.sql(
		"""select parent, supplier from `tabItem Supplier`
		    where parenttype = 'Item' and parent in %(items)s and ifnull(supplier, '') != ''
		    order by parent, idx""",
		{"items": tuple(item_codes)},
	):
		vendors.setdefault(item_code, supplier)
	return vendors


def get_lines(filters):
	# fees, FREE rows etc. are kept out by unticking Maintain Stock on the item
	conditions = ["so.docstatus = 1", "ifnull(item.is_stock_item, 0) = 1"]
	if filters.company:
		conditions.append("so.company = %(company)s")
	if filters.from_date:
		conditions.append("so.transaction_date >= %(from_date)s")
	if filters.to_date:
		conditions.append("so.transaction_date <= %(to_date)s")

	lines = frappe.db.sql(
		f"""
		select
			so.name as sales_order,
			so.custom_pick_up_location as pick_up_location,
			pl.location_number,
			ifnull(pl.is_home_delivery, 0) as is_home_delivery,
			soi.item_code,
			soi.item_name,
			soi.stock_qty as qty,
			soi.total_weight
		from `tabSales Order` so
		join `tabSales Order Item` soi on soi.parent = so.name and soi.parenttype = 'Sales Order'
		left join `tabItem` item on item.name = soi.item_code
		left join `tabPick-Up Location` pl on pl.name = so.custom_pick_up_location
		where {" and ".join(conditions)}
		order by soi.item_code, pl.location_number, so.name
		""",
		filters,
		as_dict=True,
	)

	vendors = get_vendor_map({d.item_code for d in lines})
	for d in lines:
		d.vendor = vendors.get(d.item_code, "")
	return lines


def in_view(line, view, location=None):
	if view == HOME:
		return bool(line.is_home_delivery)
	if view == PICKUP:
		if not line.pick_up_location or line.is_home_delivery:
			return False
		return not location or line.pick_up_location == location
	return True


def get_item_data(filters, view, location=None, lines=None):
	if lines is None:
		lines = get_lines(filters)
	lines = [d for d in lines if in_view(d, view, location)]

	if filters.show_order_lines:
		lines = sorted(
			lines,
			key=lambda d: (
				not d.vendor,
				d.vendor,
				d.item_code,
				d.location_number if d.location_number is not None else 999,
				d.sales_order,
			),
		)
		rows = [
			{
				"item_name": d.item_name,
				"item_code": d.item_code,
				"qty": flt(d.qty),
				"weight": flt(flt(d.total_weight) / flt(d.qty), 2) if flt(d.qty) else 0,
				"total_weight": flt(d.total_weight),
				"vendor": d.vendor,
				"sales_order": d.sales_order,
				"pick_up_location": d.pick_up_location or "",
			}
			for d in lines
		]
	else:
		rows = group_by_item(lines)

	if rows:
		rows.append(
			{
				"item_name": _("Total"),
				"qty": sum(r["qty"] for r in rows),
				"total_weight": sum(r["total_weight"] for r in rows),
				"is_total_row": 1,
			}
		)
	return rows


def group_by_item(lines):
	"""One row per item, like the workbook tabs. The orders behind each row are
	listed as "SO x qty" — the same shape as the workbook's Orders column."""
	groups = OrderedDict()
	for d in lines:
		g = groups.get(d.item_code)
		if not g:
			g = groups[d.item_code] = frappe._dict(
				item_name=d.item_name,
				item_code=d.item_code,
				vendor=d.vendor,
				qty=0,
				total_weight=0,
				orders=OrderedDict(),
				locations={},
			)
		g.qty += flt(d.qty)
		g.total_weight += flt(d.total_weight)
		g.orders[d.sales_order] = g.orders.get(d.sales_order, 0) + flt(d.qty)
		if d.pick_up_location:
			g.locations[d.pick_up_location] = d.location_number

	rows = []
	for g in groups.values():
		rows.append(
			{
				"item_name": g.item_name,
				"item_code": g.item_code,
				"qty": g.qty,
				"weight": flt(g.total_weight / g.qty, 2) if g.qty else 0,
				"total_weight": g.total_weight,
				"vendor": g.vendor,
				"sales_order": ", ".join(f"{so} x {fmt_qty(q)}" for so, q in g.orders.items()),
				"pick_up_location": ", ".join(sorted(g.locations, key=lambda n: g.locations[n] or 0)),
			}
		)
	rows.sort(key=lambda r: (not r["vendor"], r["vendor"], r["item_code"]))
	return rows


def fmt_qty(q):
	return str(int(q)) if flt(q) == int(q) else str(flt(q, 2))


def get_locations():
	"""Home Delivery first, then the pick-up locations by number — the order of
	the workbook's Totals tab."""
	return frappe.get_all(
		"Pick-Up Location",
		filters={"disabled": 0},
		fields=["name", "location_number", "is_home_delivery"],
		order_by="is_home_delivery desc, location_number asc",
	)


def get_totals_data(filters, lines=None):
	"""Per vendor: one row per location (zero rows kept, as in the workbook),
	then a Totals row. Orders with no pick-up location are not in any block."""
	if lines is None:
		lines = get_lines(filters)

	sums = defaultdict(lambda: [0.0, 0.0])
	vendors = []
	for d in lines:
		if d.vendor not in vendors:
			vendors.append(d.vendor)
		if d.pick_up_location:
			s = sums[(d.vendor, d.pick_up_location)]
			s[0] += flt(d.qty)
			s[1] += flt(d.total_weight)

	locations = [l.name for l in get_locations()]
	# a disabled location that still has orders must not drop out of the totals
	for _vendor, loc in sums:
		if loc not in locations:
			locations.append(loc)

	rows = []
	for vendor in sorted(vendors, key=lambda v: (not v, v)):
		block = []
		for loc in locations:
			qty, weight = sums.get((vendor, loc), (0.0, 0.0))
			block.append(
				{"vendor": vendor or _(NO_VENDOR), "location": loc, "qty": qty, "total_weight": weight}
			)
		rows.extend(block)
		rows.append(
			{
				"vendor": _("Totals"),
				"qty": sum(r["qty"] for r in block),
				"total_weight": sum(r["total_weight"] for r in block),
				"is_total_row": 1,
			}
		)
	return rows


# ---------------------------------------------------------------- workbook


@frappe.whitelist()
def download_workbook(filters=None):
	"""The full multi-tab workbook, same tabs as the client's spreadsheet."""
	if not frappe.get_doc("Report", "Seasonal Supplier Quantities").is_permitted():
		frappe.throw(_("Not permitted"), frappe.PermissionError)

	from openpyxl import Workbook
	from openpyxl.styles import Font

	filters = frappe._dict(frappe.parse_json(filters) if filters else {})
	lines = get_lines(filters)
	columns = get_columns()

	wb = Workbook()
	wb.remove(wb.active)
	bold = Font(bold=True)

	def add_sheet(title, cols, rows):
		ws = wb.create_sheet(title)
		ws.append([c["label"] for c in cols])
		for cell in ws[1]:
			cell.font = bold
		for r in rows:
			ws.append([r.get(c["fieldname"]) for c in cols])
			if r.get("is_total_row"):
				for cell in ws[ws.max_row]:
					cell.font = bold
		for i, c in enumerate(cols, start=1):
			ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, c["width"] // 7)
		ws.freeze_panes = "A2"

	add_sheet("All Locations", columns, get_item_data(filters, ALL, lines=lines))
	add_sheet("Pickup_Locations", columns, get_item_data(filters, PICKUP, lines=lines))
	add_sheet("Home_Deliveries", columns, get_item_data(filters, HOME, lines=lines))
	for loc in get_locations():
		if loc.is_home_delivery:
			continue
		add_sheet(sheet_title(loc.name, wb.sheetnames), columns, get_item_data(filters, PICKUP, loc.name, lines=lines))
	add_sheet("Totals", get_totals_columns(), get_totals_data(filters, lines=lines))

	out = BytesIO()
	wb.save(out)
	frappe.response["filename"] = "Seasonal Supplier Quantities.xlsx"
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


def sheet_title(name, taken):
	"""'1 - Ma's Pantry' -> '1_Mas_Pantry' (Excel: max 31 chars, no []:*?/\\)."""
	title = re.sub(r"[^\w&]+", "_", name.replace("'", "").replace(" - ", "_")).strip("_")[:31] or "Sheet"
	base, n = title, 2
	while title in taken:
		title = f"{base[:28]}_{n}"
		n += 1
	return title
